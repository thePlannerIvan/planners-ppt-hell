#!/usr/bin/env python3
"""Focused regressions for chart geometry and image fitting."""

import importlib.util
import struct
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

SCRIPT = Path(__file__).resolve().parents[1] / "native_svg_to_ppt.py"
sys.path.insert(0, str(SCRIPT.parent))
SPEC = importlib.util.spec_from_file_location("native_svg_to_ppt", SCRIPT)
CONVERTER = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = CONVERTER
SPEC.loader.exec_module(CONVERTER)

import validate_svg_layout as VALIDATOR  # noqa: E402  (shares the converter's coordinate model)


def png_bytes(width, height):
    raw = b"".join(b"\x00" + b"\x33\x77\xaa" * width for _ in range(height))

    def chunk(kind, payload):
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


class ConverterGeometryTests(unittest.TestCase):
    def test_svg_arc_uses_real_ellipse_geometry(self):
        points = CONVERTER._parse_path_to_points("M 100 100 A 50 50 0 0 1 200 100")
        arc = points[0]
        self.assertGreater(len(arc), 20)
        self.assertAlmostEqual(arc[-1][0], 200.0, places=5)
        self.assertAlmostEqual(arc[-1][1], 100.0, places=5)
        self.assertGreater(max(abs(y - 100.0) for _, y in arc), 45.0)

    def test_transform_order_and_nested_translation(self):
        self.assertEqual(CONVERTER.parse_axis_aligned_transform("translate(100 50) scale(2)"), (100.0, 50.0, 2.0, 2.0))
        self.assertEqual(CONVERTER.parse_axis_aligned_transform("scale(2) translate(100 50)"), (200.0, 100.0, 2.0, 2.0))
        self.assertEqual(CONVERTER.parse_axis_aligned_transform("translate(5 5) scale(2) translate(10 10)"), (25.0, 25.0, 2.0, 2.0))
        self.assertEqual(CONVERTER.compose_axis_aligned((5.0, 5.0, 1.0, 1.0), (20.0, 20.0, 2.0, 2.0)), (25.0, 25.0, 2.0, 2.0))

    def test_validator_asks_the_converter_for_coordinates(self):
        """检查画面不得再自己算坐标：两条路径必须给出同一个答案。

        两份实现曾经分歧到同一份输入差 100px（scale(2) translate(100 50)）。
        """
        root = etree.fromstring(
            '<svg><g transform="translate(5 5)"><g transform="scale(2) translate(10 10)">'
            '<rect x="7" y="3" width="10" height="10"/></g></g></svg>')
        parent_map = VALIDATOR.build_parent_map(root)
        rect = root.find(".//rect")
        self.assertEqual(tuple(VALIDATOR.get_accumulated_transform(rect, parent_map)), (25.0, 25.0, 2.0, 2.0))
        # the box the validator predicts is where the converter's own rule puts it
        self.assertEqual(tuple(VALIDATOR.rect_box(rect, parent_map)), (39.0, 31.0, 20.0, 20.0))
        # the historical 100px disagreement: one <g transform="scale(2) translate(100 50)">
        old = etree.fromstring(
            '<svg><g transform="scale(2) translate(100 50)"><rect x="0" y="0" width="10" height="10"/></g></svg>')
        old_map = VALIDATOR.build_parent_map(old)
        old_rect = old.find(".//rect")
        self.assertEqual(tuple(VALIDATOR.get_accumulated_transform(old_rect, old_map)), (200.0, 100.0, 2.0, 2.0))
        self.assertEqual(tuple(VALIDATOR.rect_box(old_rect, old_map)), (200.0, 100.0, 20.0, 20.0))

    def test_picture_meet_and_slice_preserve_ratio(self):
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "wide.png"
            image_path.write_bytes(png_bytes(200, 100))
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            meet = CONVERTER.add_fitted_picture(slide, str(image_path), 0, 0, 2, 2, "xMidYMid meet")
            self.assertAlmostEqual(meet.width / meet.height, 2.0, places=2)
            sliced = CONVERTER.add_fitted_picture(slide, str(image_path), 3, 0, 2, 2, "xMidYMid slice")
            self.assertEqual(sliced.width, Inches(2))
            self.assertEqual(sliced.height, Inches(2))
            self.assertGreater(sliced.crop_left, 0)
            self.assertGreater(sliced.crop_right, 0)


class ConverterStrokeWidthTests(unittest.TestCase):
    """描边保真：导出稿的描边必须与批准 PNG 同权重。

    背景：转换器曾对 stroke-width 额外乘以 0.6，使每一根描边只有声明值的 60%
    （1920 画布上 1px -> 0.30pt），细线在 PPTX 里整条消失，
    而批准稿 PNG 是按满权重渲染的，两者永远对不上。
    """

    def _line_width(self, svg_stroke_width, canvas=(1920, 1080)):
        """按真实转换路径换算一条描边，返回 PPT 里的线宽（EMU Length）。"""
        CONVERTER.set_canvas(*canvas)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
        node = etree.fromstring(
            '<rect stroke="#000000" stroke-width="%s"/>' % svg_stroke_width
        )
        CONVERTER.apply_fill_stroke(shape, node)
        return shape.line.width

    def test_stroke_width_has_no_extra_scaling(self):
        # 用 EMU 整数比较，避开 pptx 长度取整带来的浮点噪声
        for declared in (1, 2, 4):
            with self.subTest(declared=declared):
                self.assertEqual(
                    int(self._line_width(declared)),
                    int(Pt(declared * CONVERTER.FONT_SCALE)),
                    f"{declared}px 描边被额外缩放了",
                )

    def test_one_px_stroke_is_not_flagged(self):
        # 1920 画布上的 1px ≈ 0.49999pt，是合法的发丝线，不应触发告警
        CONVERTER._conversion_warnings.clear()
        self._line_width(1)
        self.assertEqual(
            [w for w in CONVERTER._conversion_warnings if "stroke-width" in w],
            [],
            "1px 描边被误判为过细",
        )

    def test_sub_half_point_stroke_is_reported(self):
        CONVERTER._conversion_warnings.clear()
        self._line_width(0.5)
        self.assertTrue(
            any("stroke-width" in w for w in CONVERTER._conversion_warnings),
            "低于 0.5pt 的描边必须给出提示，否则细线会静默消失",
        )


if __name__ == "__main__":
    unittest.main()
