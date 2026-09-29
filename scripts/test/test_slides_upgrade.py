#!/usr/bin/env python3
"""
Regression suite for the PPT (slides) route upgrades:
1. `native_svg_to_ppt.py` (v6.0) — 6 deterministic upgrades (`<style>` + `:root` `var(--...)`,
   `<tspan dy>`, leaf-node `transform` composition, `1px` `<rect>` divider preservation,
   `<text>`/`<g>` `opacity` + `letter-spacing * FONT_SCALE` + `dominant-baseline` + `rgba()`,
   and `<linearGradient>` -> `<a:gradFill>`).
2. `validate_svg_layout.py` + `project_state.py` (v6.0) — 3 new hard guards (`LINE_CROSSES_TEXT`,
   `EMOJI_IN_SLIDE_TEXT`, `TEXT_OVERFLOWS_CONTAINER`) + `design_observations` leaf transform / rect bounds.
3. Multimodal 4-Piece Template Packs (`agency-social-proposal`, `consulting-product-strategy`,
   `extract_template_pack.py`, `template_library.py` v2 + v1 compatibility).
4. Simplified PPT-Style Visual Review Workbench (`review.html`, `generate_review_html.py`, `review_feedback.py`
   with `data-review-id` injection and deterministic `svg_edits` write-back on `consume(root)`).
"""

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from pptx import Presentation

SCRIPTS = Path(__file__).resolve().parents[1]
SKILL_ROOT = SCRIPTS.parent
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(SCRIPTS / "template"))
sys.path.insert(0, str(SCRIPTS / "orchestrate"))
sys.path.insert(0, str(SCRIPTS / "test"))

import native_svg_to_ppt as conv  # noqa: E402
import validate_svg_layout as val  # noqa: E402
import project_state as ps  # noqa: E402
from template_library import LIBRARY_ROOT, apply_template, list_templates, validate_pack_dir  # noqa: E402
import test_v5 as fixtures  # noqa: E402
from ppt_pipeline import make_review, next_action, unresolved  # noqa: E402
from review_feedback import consume  # noqa: E402


TORTURE_SVG = """<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" width="1920" height="1080">
  <defs>
    <linearGradient id="heroGrad" x1="0%" y1="0%" x2="100%" y2="0%">
      <stop offset="0%" stop-color="#0F172A" stop-opacity="1"/>
      <stop offset="100%" stop-color="#0052D9" stop-opacity="0.85"/>
    </linearGradient>
  </defs>
  <style>
    :root {
      --bg-canvas: #F4F5F7;
      --brand-red: #E60012;
      --ink-dark: #111111;
    }
    .canvas-bg { fill: var(--bg-canvas); }
    .card-title { fill: var(--ink-dark); font-size: 28px; font-weight: 700; font-family: "PingFang SC", sans-serif; }
    #accent-bar { fill: var(--brand-red); }
  </style>

  <rect class="canvas-bg" x="0" y="0" width="1920" height="1080"/>
  <rect id="grad-banner" x="80" y="64" width="1760" height="96" rx="12" fill="url(#heroGrad)"/>
  <rect id="thin-divider-h" x="80" y="180" width="1760" height="1" fill="#CBD5E1"/>
  <rect id="thin-divider-v" x="640" y="200" width="1" height="680" fill="#CBD5E1"/>

  <g transform="translate(80, 220)">
    <rect id="leaf-tx-rect" x="0" y="0" width="520" height="320" rx="12" fill="#FFFFFF" transform="translate(20, 30)"/>
    <text id="watermark-num" x="380" y="140" font-family="JetBrains Mono, monospace" font-size="140" font-weight="900" fill="#111111" opacity="0.12">01</text>
    <g opacity="0.25">
      <text id="group-opacity-text" x="40" y="80" font-family="PingFang SC" font-size="24" fill="#E60012">GROUP_OPACITY</text>
    </g>
    <rect x="40" y="110" width="220" height="40" rx="20" fill="rgba(230, 0, 18, 0.15)"/>
    <text id="pill-text" x="150" y="130" text-anchor="middle" dominant-baseline="central"
          font-family="PingFang SC" font-size="16" font-weight="700" letter-spacing="4px" fill="rgba(230, 0, 18, 0.85)">核心战略胶囊</text>

    <text id="multiline-text" class="card-title" x="0" y="0" transform="translate(40, 200)">
      <tspan x="0" dy="0">第一行主判断：</tspan><tspan fill="var(--brand-red)">核心高亮词</tspan>
      <tspan x="0" dy="36" font-size="18" font-weight="400" fill="#4B5563">第二行详细论据说明文字（独立段落）</tspan>
      <tspan x="0" dy="28" font-size="18" font-weight="400" fill="#4B5563">第三行补充口径说明文字（独立段落）</tspan>
    </text>
  </g>

  <line x1="80" y1="900" x2="1840" y2="900" stroke="var(--brand-red)" stroke-width="2" stroke-dasharray="8 6"/>
  <rect id="accent-bar" x="80" y="920" width="1760" height="76" rx="12"/>
</svg>
"""


def convert_svgs_to_pptx(svg_paths, output_pptx):
    env = os.environ.copy()
    env["SMART_SVG_EXPORT_APPROVED_BY_PIPELINE"] = "1"
    cmd = [
        sys.executable,
        str(SCRIPTS / "native_svg_to_ppt.py"),
        *[str(p) for p in svg_paths],
        "-o",
        str(output_pptx),
    ]
    res = subprocess.run(cmd, env=env, capture_output=True, text=True)
    if res.returncode != 0:
        raise AssertionError(f"native_svg_to_ppt.py failed ({res.returncode}):\n{res.stdout}\n{res.stderr}")
    return res


class ConverterAndValidatorUpgradeTests(unittest.TestCase):
    def test_converter_six_upgrades_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            svg_path = Path(tmp) / "slide_1.svg"
            pptx_path = Path(tmp) / "out.pptx"
            svg_path.write_text(TORTURE_SVG, encoding="utf-8")

            res = val.validate_file(svg_path, route="slides")
            errors = [i for i in res["issues"] if i["severity"] == "error"]
            warnings = [i for i in res["issues"] if i["severity"] == "warning"]
            self.assertEqual(errors, [], f"Unexpected errors on TORTURE_SVG: {errors}")
            self.assertEqual(warnings, [], f"Unexpected warnings on TORTURE_SVG: {warnings}")

            convert_svgs_to_pptx([svg_path], pptx_path)
            prs_loaded = Presentation(pptx_path)
            slide = prs_loaded.slides[0]
            shapes = list(slide.shapes)

            # 1 & 6: Native <a:gradFill>
            grad_fills = slide._element.xpath(".//a:gradFill")
            self.assertGreaterEqual(len(grad_fills), 1)
            self.assertEqual(len(grad_fills[0].xpath(".//a:gs")), 2)

            # 4: 1px thin divider preservation
            one_px_in = conv.svg_to_inches(1.0)
            thin_shapes = [
                s for s in shapes
                if abs(s.height.inches - one_px_in) < 0.002 or abs(s.width.inches - one_px_in) < 0.002
            ]
            self.assertGreaterEqual(len(thin_shapes), 2)

            # 3: Leaf-node transform composition
            expected_left = conv.svg_to_inches(100.0)
            expected_top = conv.svg_to_inches(250.0)
            expected_w = conv.svg_to_inches(520.0)
            leaf_rects = [
                s for s in shapes
                if abs(s.left.inches - expected_left) < 0.02
                and abs(s.top.inches - expected_top) < 0.02
                and abs(s.width.inches - expected_w) < 0.02
            ]
            self.assertEqual(len(leaf_rects), 1)

            # 2: <tspan dy> multi-paragraph conversion
            text_shapes = [s for s in shapes if s.has_text_frame]
            ml_shape = [s for s in text_shapes if "第一行主判断" in s.text_frame.text][0]
            self.assertEqual(len(ml_shape.text_frame.paragraphs), 3)
            self.assertGreaterEqual(len(ml_shape.text_frame.paragraphs[0].runs), 2)
            self.assertAlmostEqual(ml_shape.left.inches, conv.svg_to_inches(120.0), delta=0.05)

            # 5: <text opacity>, <g opacity>, rgba(), letter-spacing, dominant-baseline
            wm_shape = [s for s in text_shapes if s.text_frame.text.strip() == "01"][0]
            self.assertEqual(wm_shape._element.xpath(".//a:rPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["12000"])
            grp_shape = [s for s in text_shapes if s.text_frame.text.strip() == "GROUP_OPACITY"][0]
            self.assertEqual(grp_shape._element.xpath(".//a:rPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["25000"])
            pill_shape = [s for s in text_shapes if s.text_frame.text.strip() == "核心战略胶囊"][0]
            self.assertEqual(pill_shape._element.xpath(".//a:rPr/@spc"), ["200"])
            self.assertEqual(pill_shape._element.xpath(".//a:rPr/a:solidFill/a:srgbClr/a:alpha/@val"), ["85000"])
            self.assertAlmostEqual(pill_shape.top.inches, conv.svg_to_inches(341.36), delta=0.03)

    def test_validator_new_hard_guards_synthetic(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "bad.svg"
            p.write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" width="1920" height="1080">'
                '<rect x="0" y="0" width="1920" height="1080" fill="#FFFFFF"/>'
                '<rect x="100" y="100" width="120" height="36" rx="18" fill="#E60012"/>'
                '<text x="160" y="124" text-anchor="middle" font-size="22" font-family="PingFang SC" fill="#FFFFFF">'
                '这是一段远远超出120像素胶囊宽度的超长标签文字</text>'
                '<text x="100" y="300" font-size="28" font-family="PingFang SC" fill="#111111">🎯 目标人群定位</text>'
                '<text x="100" y="500" font-size="28" font-family="PingFang SC" fill="#111111">'
                '<tspan x="100" dy="0">第一行文字</tspan>'
                '<tspan x="100" dy="34">第二行被横线穿过的文字</tspan>'
                '</text>'
                '<line x1="80" y1="530" x2="600" y2="530" stroke="#CBD5E1" stroke-width="2" stroke-dasharray="6 4"/>'
                '</svg>',
                encoding="utf-8",
            )
            res = val.validate_file(p, route="slides")
            codes = {i["code"]: i["severity"] for i in res["issues"]}
            self.assertEqual(codes.get("EMOJI_IN_SLIDE_TEXT"), "error")
            self.assertEqual(codes.get("LINE_CROSSES_TEXT"), "error")
            self.assertEqual(codes.get("TEXT_OVERFLOWS_CONTAINER"), "warning")

    def test_design_observations_leaf_transforms_and_rect_bounds(self):
        with tempfile.TemporaryDirectory() as tmp:
            root_dir = Path(tmp)
            svg_dir = root_dir / ps.SVG
            svg_dir.mkdir(parents=True)
            (svg_dir / "slide_1.svg").write_text(TORTURE_SVG, encoding="utf-8")
            obs = ps.design_observations(root_dir, "slide_1")
            self.assertAlmostEqual(obs["lowest_content_bottom"], 996.0, delta=1.0)
            self.assertIn(140, obs["type_tiers"])
            self.assertIn(16, obs["type_tiers"])


class MultimodalTemplatePackTests(unittest.TestCase):
    def test_builtin_flagship_template_packs_are_valid_and_pass_slide_validator(self):
        items = {item["template_id"]: item for item in list_templates()}
        self.assertIn("agency-social-proposal", items)
        self.assertIn("consulting-product-strategy", items)

        for pack_id in ("agency-social-proposal", "consulting-product-strategy"):
            pack_dir = LIBRARY_ROOT / pack_id
            v = validate_pack_dir(pack_dir, check_manifest_hashes=True)
            self.assertTrue(v["valid"], f"{pack_id} failed validation: {v['issues']}")
            self.assertEqual(v["primitive_count"], 6)
            for svg_path in sorted((pack_dir / "primitives").glob("*.svg")):
                res = val.validate_file(svg_path, route="slides")
                errors = [i for i in res["issues"] if i["severity"] == "error"]
                warnings = [i for i in res["issues"] if i["severity"] == "warning"]
                self.assertEqual(errors, [], f"{pack_id}/{svg_path.name} errors: {errors}")
                self.assertEqual(warnings, [], f"{pack_id}/{svg_path.name} warnings: {warnings}")

    def test_apply_v2_template_pack_into_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            proj = Path(tmp) / "demo_proj"
            out = apply_template(proj, "agency-social-proposal")
            self.assertEqual(out["applied_template_id"], "agency-social-proposal")
            self.assertTrue((proj / "_internal/00_project/template_pack/tokens.css").is_file())
            self.assertTrue((proj / "_internal/00_project/template_pack/SPEC.md").is_file())


class ReviewWorkbenchInlineEditsTests(unittest.TestCase):
    def test_svg_edits_writeback_and_delete_on_consume(self):
        fixture = fixtures.V5Tests()
        fixture.setUp()
        try:
            root = fixture.root.resolve()
            (root / ps.SVG / "alpha.svg").write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" viewBox="0 0 1920 1080">'
                '<rect width="1920" height="1080" fill="#FFFFFF"/>'
                '<rect x="120" y="260" width="480" height="220" rx="12" fill="#E2E8F0"/>'
                '<text x="120" y="180" font-size="48" font-family="Arial" fill="#111111">Evidence</text>'
                '</svg>',
                encoding="utf-8",
            )
            fixture.seal()
            snap = make_review(root)
            payload = {
                "review_id": snap["review_id"],
                "pages": {
                    "alpha": {
                        "decision": "approved",
                        "feedback": "",
                        "annotations": [],
                        "assets": [],
                        "svg_edits": [
                            {"review_id": "0.1", "tag": "rect", "dx": 0, "dy": 0, "deleted": True},
                            {"review_id": "0.2", "tag": "text", "text": "Direct Title", "font_size": 52, "dx": 24, "dy": 12},
                        ],
                    },
                    "omega": {
                        "decision": "approved",
                        "feedback": "",
                        "annotations": [],
                        "assets": [],
                    },
                },
                "overall_feedback": "",
            }
            intake = fixture.submit(payload)
            self.assertFalse(intake["stale"])
            self.assertIn("alpha", intake["svg_edits_applied"])
            alpha_svg = (root / ps.SVG / "alpha.svg").read_text(encoding="utf-8")
            self.assertNotIn("#E2E8F0", alpha_svg, "Deleted rect 0.1 must be removed from alpha.svg")
            self.assertIn("Direct Title", alpha_svg)
            self.assertIn('font-size="52"', alpha_svg)
            self.assertEqual(unresolved(root), [])
            self.assertTrue(ps.approved(root))
            self.assertEqual(next_action(root)["state"], "EXPORT")
        finally:
            fixture.tearDown()


if __name__ == "__main__":
    unittest.main()
