"""Candidate pipeline regressions; synthetic previews are protocol fixtures, not visual review."""
import json
import shutil
import sys
import tempfile
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

SCRIPTS=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(SCRIPTS))
sys.path.insert(0,str(SCRIPTS/'orchestrate'))
sys.path.insert(0,str(SCRIPTS/'template'))
import ppt_pipeline as pipeline
import project_state as ps
import review_feedback
import template_feedback
import template_library
import workbench_store as store
from generate_template_review_html import generate_pack_review


def svg(text='Original'):
    return ('<svg xmlns="http://www.w3.org/2000/svg" width="1920" height="1080" '
            'viewBox="0 0 1920 1080"><rect width="1920" height="1080" fill="#fff"/>'
            '<text id="title" x="120" y="180" font-size="48" fill="#111111" font-family="Arial">'
            +text+'</text></svg>')


class WorkbenchPipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)/'project'
        for directory in (ps.PROJECT,ps.SVG,ps.VALIDATION,ps.REVIEW,ps.PNG):
            (self.root/directory).mkdir(parents=True,exist_ok=True)
        ps.write(self.root/ps.PROJECT/'page_manifest.json',
                 {'version':'5.0','route':'slides','template_intake':{'mode':'autonomous'}})
        ps.write(self.root/ps.CONTENT,{'project':'Test','pages':[
            {'page_key':key,'title':key,'content':'Original','notes':'notes','source_assets':[]}
            for key in ('alpha','omega')]})
        (self.root/ps.DIRECTION).write_text('# Direction\n',encoding='utf-8')
        for key in ('alpha','omega'):
            (self.root/ps.SVG/(key+'.svg')).write_text(svg(),encoding='utf-8')
        self.sequence=0
        store.ensure(self.root)

    def tearDown(self): self.tmp.cleanup()

    def command(self,**fields):
        self.sequence+=1
        return {'operation_id':'operation_'+str(self.sequence),**fields}

    def save(self,key='alpha',text='New',notes=None):
        checked=store.commit(self.root,self.command(op='checkout',page_key=key))
        self.assertTrue(checked['ok'])
        candidate=Path(checked['candidate'])
        candidate.write_text(candidate.read_text().replace('Original',text),encoding='utf-8')
        cmd=self.command(op='save',author='model',page_key=key,candidate=str(candidate),
                         base_revision=checked['base_revision'])
        if notes is not None: cmd['notes']=notes
        saved=store.commit(self.root,cmd)
        self.assertTrue(saved['ok'],saved)
        return saved

    def feedback(self):
        revision=store.get_page(self.root,'alpha')['revision']
        result=store.commit(self.root,self.command(op='feedback',scope='page',
                    pages={'alpha':{'revision':revision,'feedback':'Improve title'}}),browser=True)
        self.assertTrue(result['ok'],result)
        return result['task_id']

    def fake_render(self,*args,**kwargs):
        directory=Path(args[2]);source=Path(args[3])
        Image.new('RGB',(192,108),'white').save(directory/(source.stem+'.png'))
        return ''

    def test_pending_tasks_route_to_checkout_without_legacy_consumption(self):
        task=self.feedback()
        raw={'pages':{'alpha':{'svg_edits':[{'review_id':'0.1','text':'Wrong'}]}}}
        ps.write(self.root/ps.REVIEW/'feedback.json',raw)
        before=(self.root/ps.SVG/'alpha.svg').read_bytes()
        action=pipeline.next_action(self.root)
        self.assertEqual(action['state'],'WORKBENCH')
        self.assertEqual(action['feedback'][0]['task_id'],task)
        self.assertIn('checkout',action['instruction'])
        self.assertEqual(action['legacy_feedback_recovery']['raw_pending'],raw)
        self.assertEqual(ps.read(self.root/ps.REVIEW/'feedback.json'),raw)
        self.assertEqual((self.root/ps.SVG/'alpha.svg').read_bytes(),before)

    def test_legacy_apply_disabled_even_if_called_directly(self):
        with self.assertRaisesRegex(ValueError,'Legacy tree-index'):
            review_feedback.apply_pending_svg_edits(self.root,{'pages':{}})

    def test_empty_page_next_uses_candidate_not_compat_overwrite(self):
        other=Path(self.tmp.name)/'empty'
        shutil.copytree(self.root/ps.PROJECT,other/ps.PROJECT)
        ps.write(other/ps.CONTENT,ps.read(self.root/ps.CONTENT))
        action=pipeline.next_action(other)
        self.assertEqual(action['state'],'WORKBENCH')
        self.assertIn('checkout',action['instruction'])
        self.assertIn('base_revision',action['instruction'])

    def test_external_compat_change_is_reported_not_imported(self):
        revision=store.get_page(self.root,'alpha')['revision']
        (self.root/ps.SVG/'alpha.svg').write_text(svg('External'))
        action=pipeline.next_action(self.root)
        self.assertEqual(action['state'],'WORKBENCH')
        self.assertIn('outside workbench',action['error'])
        self.assertEqual(ps.read(self.root/ps.WORKBENCH/'head.json')['pages']['alpha']['revision'],revision)

    def test_saved_notes_invalidate_page_version(self):
        before=ps.page_version(self.root,ps.content(self.root)['pages'][0])
        self.save(text='Original',notes='Updated notes only')
        after=ps.page_version(self.root,ps.content(self.root)['pages'][0])
        self.assertNotEqual(before,after)
        self.assertEqual(ps.versions(self.root)['alpha'],after)

    def test_input_baseline_invalidation_agrees_between_check_and_versions(self):
        before=ps.versions(self.root)
        contents=ps.content(self.root)
        contents['pages'][0]['content']='New evidence baseline'
        ps.write(self.root/ps.CONTENT,contents)
        after=ps.versions(self.root)
        self.assertNotEqual(before['alpha'],after['alpha'])
        self.assertEqual(before['omega'],after['omega'])

    def test_resolve_binds_saved_revision(self):
        task=self.feedback();result=self.save()
        receipt=pipeline.resolve(self.root,task,'Changed headline',result['revision'],'resolve_one')
        self.assertTrue(receipt['ok'],receipt)
        self.assertEqual(store.state(self.root)['tasks'][task]['results'],{'alpha':result['revision']})
        self.assertEqual(pipeline.unresolved(self.root),[])

    def test_export_without_approval_and_pending_feedback_uses_frozen_order(self):
        task=self.feedback();head=store.state(self.root)
        order=store.commit(self.root,self.command(op='order',order=['omega','alpha'],base_order=head['order']))
        self.assertTrue(order['ok'])
        record=pipeline.export(self.root)
        self.assertEqual(record['order'],['omega','alpha'])
        frozen=ps.read(record['snapshot_path'])
        self.assertEqual(frozen['order'],record['order'])
        self.assertEqual(store.state(self.root)['tasks'][task]['status'],'pending')
        self.assertEqual(len(Presentation(self.root/'final_deck.pptx').slides),2)
        self.assertNotIn('reviewed',record)

    def test_export_live_edit_during_conversion_does_not_change_snapshot(self):
        real_run=pipeline.run
        def convert(*args,**kwargs):
            if str(args[0]).endswith('native_svg_to_ppt.py'):
                self.assertIn('/snapshots/',str(args[1]))
                self.assertIn('/scenes/',str(args[1]))
                self.save(text='Later')
            return real_run(*args,**kwargs)
        with patch.object(pipeline,'run',side_effect=convert):
            record=pipeline.export(self.root)
        prs=Presentation(self.root/'final_deck.pptx')
        texts=[shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertIn('Original',texts);self.assertNotIn('Later',texts)
        frozen=ps.read(record['snapshot_path'])
        self.assertNotEqual(frozen['pages'][0]['revision'],store.get_page(self.root,'alpha')['revision'])

    def test_selected_snapshot_does_not_consult_live_order_or_content(self):
        frozen=store.snapshot(self.root)
        self.save(text='Later')
        head=store.state(self.root)
        store.commit(self.root,self.command(op='order',order=['omega','alpha'],base_order=head['order']))
        # A selected snapshot remains usable even if the live content draft is incomplete.
        ps.write(self.root/ps.CONTENT,{'pages':[]})
        record=pipeline.export(self.root,snapshot_path=frozen['path'])
        self.assertEqual(record['snapshot_id'],frozen['snapshot_id'])
        self.assertEqual(record['order'],['alpha','omega'])
        texts=[shape.text for slide in Presentation(record['output_path']).slides
               for shape in slide.shapes if shape.has_text_frame]
        self.assertIn('Original',texts);self.assertNotIn('Later',texts)

    def test_invalid_selected_snapshot_never_falls_back(self):
        frozen=store.snapshot(self.root)
        Path(frozen['root'],'scenes','alpha.svg').write_text(svg('Tampered'))
        with self.assertRaises(ValueError): pipeline.export(self.root,snapshot_path=frozen['path'])
        self.assertFalse((self.root/'final_deck.pptx').exists())

    def test_failed_validator_cannot_export_using_a_zero_error_summary(self):
        frozen=store.snapshot(self.root)
        failed=subprocess.CompletedProcess([],2,stdout='{"summary":{"errors":0}}',stderr='Validator failed')
        with patch.object(pipeline.subprocess,'run',return_value=failed):
            with self.assertRaisesRegex(ValueError,'Validator failed'):
                pipeline.export(self.root,snapshot_path=frozen['path'])
        self.assertFalse((self.root/'final_deck.pptx').exists())

    def test_job_isolated_notes_report_and_cached_snapshot_output(self):
        frozen=store.snapshot(self.root)
        first=pipeline.export(self.root,snapshot_path=frozen['path'])
        job=self.root.resolve()/pipeline.OUTPUT/frozen['snapshot_id']
        self.assertEqual(Path(first['output_path']),job/'final_deck.pptx')
        self.assertEqual(Path(first['notes_path']).parent,job)
        self.assertEqual(Path(first['conversion_report']).parent,job)
        with patch.object(pipeline,'run',side_effect=AssertionError('Must use cached output')):
            again=pipeline.export(self.root,snapshot_path=frozen['path'])
        self.assertEqual(first,again)

    def concurrent_exports(self,snapshots):
        processes=[subprocess.Popen([sys.executable,str(SCRIPTS/'orchestrate/ppt_pipeline.py'),
                    str(self.root),'export','--snapshot',frozen['path']],
                    stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True) for frozen in snapshots]
        results=[]
        try:
            for process in processes:
                stdout,stderr=process.communicate(timeout=45)
                self.assertEqual(process.returncode,0,stderr or stdout)
                results.append(json.loads(stdout))
        finally:
            for process in processes:
                if process.poll() is None:
                    process.kill();process.communicate()
        return results

    def test_concurrent_same_snapshot_exports_reuse_one_complete_job(self):
        frozen=store.snapshot(self.root)
        first,second=self.concurrent_exports([frozen,frozen])
        self.assertEqual(first,second)
        self.assertEqual(ps.sha(first['output_path']),first['pptx_sha256'])
        self.assertEqual(ps.sha(first['notes_path']),first['notes_sha256'])
        self.assertEqual(ps.sha(first['conversion_report']),first['conversion_report_sha256'])

    def test_concurrent_different_snapshots_keep_notes_and_outputs_separate(self):
        old=store.snapshot(self.root)
        self.save(text='Later',notes='Later notes')
        new=store.snapshot(self.root)
        first,second=self.concurrent_exports([old,new])
        self.assertNotEqual(first['output_path'],second['output_path'])
        self.assertEqual(ps.read(first['notes_path'])['alpha.svg'],'notes')
        self.assertEqual(ps.read(second['notes_path'])['alpha.svg'],'Later notes')
        for record,expected in ((first,'Original'),(second,'Later')):
            texts=[shape.text for slide in Presentation(record['output_path']).slides
                   for shape in slide.shapes if shape.has_text_frame]
            self.assertIn(expected,texts)
            self.assertEqual(ps.sha(record['output_path']),record['pptx_sha256'])
        latest=ps.read(self.root/ps.PROJECT/'export.json')
        self.assertEqual(ps.sha(self.root/'final_deck.pptx'),latest['pptx_sha256'])

    def test_cached_export_does_not_hide_damaged_notes_or_report(self):
        for field in ('notes_path','conversion_report'):
            with self.subTest(field=field):
                frozen=store.snapshot(self.root)
                record=pipeline.export(self.root,snapshot_path=frozen['path'])
                Path(record[field]).write_text('{}',encoding='utf-8')
                with self.assertRaisesRegex(ValueError,'Recorded export is damaged'):
                    pipeline.export(self.root,snapshot_path=frozen['path'])
                self.assertEqual(ps.sha(record['output_path']),record['pptx_sha256'])

    def test_export_inspect_explicit_job_survives_latest_alias_replacement(self):
        frozen=store.snapshot(self.root);old=pipeline.export(self.root,snapshot_path=frozen['path'])
        self.save(text='Later');new=pipeline.export(self.root)
        previews=[]
        for key in ('alpha','omega'):
            path=self.root/(key+'.png');Image.new('RGB',(192,108),'white').save(path)
            previews.append(path.name)
        result=pipeline.export_inspect(self.root,'Synthetic protocol fixture',previews,snapshot_path=frozen['path'])
        self.assertEqual(result['snapshot_id'],old['snapshot_id'])
        self.assertEqual(ps.read(self.root/ps.PROJECT/'export.json')['snapshot_id'],new['snapshot_id'])
        self.assertEqual(ps.sha(old['output_path']),old['pptx_sha256'])

    def test_converter_help_and_direct_use_need_no_approval_environment(self):
        import os
        env=dict(os.environ);env.pop('SMART_SVG_EXPORT_APPROVED_BY_PIPELINE',None)
        help_result=subprocess.run([sys.executable,str(SCRIPTS/'native_svg_to_ppt.py'),'--help'],
                                   capture_output=True,text=True,env=env)
        self.assertEqual(help_result.returncode,0,help_result.stderr)
        output=self.root/'direct.pptx'
        result=subprocess.run([sys.executable,str(SCRIPTS/'native_svg_to_ppt.py'),
                    str(self.root/ps.SVG/'alpha.svg'),'-o',str(output),'--strict-missing-images'],
                    capture_output=True,text=True,env=env)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertTrue(output.is_file())

    def test_export_inspection_is_bound_to_output_not_live_pages(self):
        pipeline.export(self.root);self.save(text='Later')
        previews=[]
        for key in ('alpha','omega'):
            path=self.root/(key+'.png');Image.new('RGB',(192,108),'white').save(path)
            previews.append(str(path.relative_to(self.root)))
        result=pipeline.export_inspect(self.root,'Synthetic protocol fixture',previews)
        self.assertEqual(result['visual_verification'],'inspected')
        self.assertNotEqual(pipeline.next_action(self.root)['state'],'COMPLETE')

    def test_video_export_is_snapshot_without_review(self):
        # Change route before initializing a distinct store.
        other=Path(self.tmp.name)/'video'
        shutil.copytree(self.root,other,ignore=shutil.ignore_patterns('06_workbench'))
        script=other/'script.md';script.write_text('Original narration\n')
        plan=other/'visual-plan.json'
        ps.write(plan,{'contract_version':'5.0','script_hash':ps.sha(script),'screens':[]})
        m=ps.manifest(other);m['route']='video'
        m['visual_contract']={'path':str(plan.resolve()),'sha256':ps.sha(plan),
                              'script_hash':ps.sha(script),'contract_version':'5.0'}
        ps.write(other/ps.PROJECT/'page_manifest.json',m)
        result=pipeline.export(other)
        self.assertEqual(result['route'],'video')
        self.assertEqual(ps.read(result['snapshot_path'])['route'],'video')
        self.assertFalse((other/'final_deck.pptx').exists())

    def test_missing_image_save_or_export_is_technical_error(self):
        checked=store.commit(self.root,self.command(op='checkout',page_key='alpha'))
        Path(checked['candidate']).write_text(svg().replace('</svg>','<image href="missing.png" width="100" height="100"/></svg>'))
        result=store.commit(self.root,self.command(op='save',author='model',page_key='alpha',
                    candidate=checked['candidate'],base_revision=checked['base_revision']))
        self.assertFalse(result['ok'])

    def test_illegal_svg_features_still_block_export(self):
        checked=store.commit(self.root,self.command(op='checkout',page_key='alpha'))
        candidate=Path(checked['candidate'])
        candidate.write_text(candidate.read_text().replace('fill="#111111"',''))
        result=store.commit(self.root,self.command(op='save',author='model',page_key='alpha',
                    candidate=str(candidate),base_revision=checked['base_revision']))
        self.assertTrue(result['ok'],result)
        with self.assertRaises(ValueError): pipeline.export(self.root)

    def test_render_revision_temp_then_record_with_renderer_digest(self):
        with patch.object(pipeline,'run',side_effect=self.fake_render):
            result=pipeline.check(self.root,['alpha'])
        self.assertEqual(result['pages'][0]['status'],'pass',result)
        revision=store.get_page(self.root,'alpha')['revision']
        record=ps.read(self.root/ps.WORKBENCH/'renders'/revision/'render.json')
        self.assertIn('sha256:',record['renderer'])
        self.assertEqual(ps.read(self.root/ps.VALIDATION/'alpha.json')['revision'],revision)
        self.assertTrue((self.root/ps.PNG/'alpha.png').is_file())

    def test_renderer_failure_preserves_existing_stamp_and_png(self):
        with patch.object(pipeline,'run',side_effect=self.fake_render): pipeline.check(self.root,['alpha'])
        old_stamp=(self.root/ps.VALIDATION/'alpha.json').read_bytes()
        old_png=(self.root/ps.PNG/'alpha.png').read_bytes()
        self.save()
        with patch.object(pipeline,'run',side_effect=ValueError('Renderer failed')):
            result=pipeline.check(self.root,['alpha'])
        self.assertEqual(result['pages'][0]['status'],'blocked')
        self.assertEqual((self.root/ps.VALIDATION/'alpha.json').read_bytes(),old_stamp)
        self.assertEqual((self.root/ps.PNG/'alpha.png').read_bytes(),old_png)

    def test_late_render_cannot_replace_new_png_or_stamp(self):
        previous=store.get_page(self.root,'alpha')['revision']
        sentinel={'version':'newer stamp','errors':0}
        def render_then_edit(*args,**kwargs):
            self.fake_render(*args,**kwargs)
            newer=self.save()
            image=self.root/'new.png';Image.new('RGB',(192,108),'red').save(image)
            store.record_render(self.root,'alpha',newer['revision'],image,'test:new')
            ps.write(self.root/ps.VALIDATION/'alpha.json',sentinel)
        with patch.object(pipeline,'run',side_effect=render_then_edit):
            result=pipeline.check(self.root,['alpha'])
        self.assertEqual(result['pages'][0]['status'],'blocked',result)
        self.assertEqual(ps.read(self.root/ps.VALIDATION/'alpha.json'),sentinel)
        with Image.open(self.root/ps.PNG/'alpha.png') as image:
            self.assertEqual(image.getpixel((0,0)),(255,0,0))
        self.assertTrue((self.root/ps.WORKBENCH/'renders'/previous/'page.png').is_file())


class FourPieceApprovalTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)/'project'
        self.pack=self.root/ps.PROJECT/'template_pack'
        shutil.copytree(template_library.LIBRARY_ROOT/'agency-social-proposal',self.pack)
        self.library=Path(self.tmp.name)/'library'
        ps.write(self.root/ps.PROJECT/'page_manifest.json',{'version':'5.0','route':'slides'})

    def tearDown(self): self.tmp.cleanup()

    def approve(self):
        generate_pack_review(self.root)
        snapshot=ps.template_review_snapshot(self.root)
        data={'review_id':snapshot['review_id'],'submission_action':'approve_all',
              'template_name':'Test pack','overall_feedback':'',
              'layouts':{key:{'decision':'pass','approved':True,'custom_feedback':''}
                         for key in snapshot['layouts']},
              'provenance':{'source':'template_review_page'}}
        ps.write(self.root/ps.PROJECT/'template_feedback.json',data)
        result=template_feedback.consume(self.root)
        self.assertTrue(result['document']['approved'],result)
        return result

    def test_no_approval_no_publish_or_directory_deletion(self):
        existing=self.library/'same';existing.mkdir(parents=True)
        (existing/'keep.txt').write_text('Original library package')
        with self.assertRaises(ValueError):
            template_library.publish(self.root,template_id='same',library_root=self.library,replace=True)
        self.assertTrue((existing/'keep.txt').is_file())

    def test_generation_review_consume_publish_share_four_piece_protocol(self):
        self.approve()
        expected=['skyline_shell.svg',*[path.relative_to(self.pack).as_posix()
                                       for path in sorted((self.pack/'primitives').glob('*.svg'))]]
        self.assertEqual(template_feedback.expected_layouts(self.root),expected)
        result=template_library.publish(self.root,template_id='test',library_root=self.library)
        self.assertEqual(result['status'],'approved')
        self.assertTrue(template_library.validate_pack_dir(self.library/'test',True)['valid'])

    def test_changed_tokens_cannot_reuse_approval(self):
        self.approve()
        with (self.pack/'tokens.css').open('a') as handle: handle.write('\n/* changed */\n')
        result=template_feedback.consume(self.root)
        self.assertTrue(result['stale']);self.assertFalse(result['document']['approved'])
        with self.assertRaises(ValueError): template_library.publish(self.root,library_root=self.library)

    def test_partial_or_revision_decision_does_not_publish(self):
        self.approve()
        feedback=ps.read(self.root/ps.PROJECT/'template_feedback.json')
        key=next(iter(feedback['layouts']))
        feedback['layouts'][key]={'decision':'revise','approved':False,'custom_feedback':'Change title'}
        feedback['submission_action']='submit_batch'
        ps.write(self.root/ps.PROJECT/'template_feedback.json',feedback)
        with self.assertRaises(ValueError): template_library.publish(self.root,library_root=self.library)

    def test_same_id_requires_explicit_replace_and_preserves_old_package(self):
        self.approve()
        old=template_library.publish(self.root,template_id='same',library_root=self.library)
        with self.assertRaisesRegex(ValueError,'explicit --replace'):
            template_library.publish(self.root,template_id='same',library_root=self.library)
        (self.pack/'SPEC.md').write_text('# Updated theme\n')
        self.approve()
        new=template_library.publish(self.root,template_id='same',library_root=self.library,replace=True)
        backups=list((self.library/'.history'/'same').iterdir())
        self.assertEqual(len(backups),1)
        self.assertEqual(ps.read(backups[0]/'manifest.json')['package_sha256'],old['package_sha256'])
        self.assertNotEqual(old['package_sha256'],new['package_sha256'])

    def test_manifest_builder_does_not_grant_approval(self):
        manifest=template_library.build_manifest(self.pack,'test','Test')
        self.assertEqual(manifest['status'],'candidate')


if __name__=='__main__': unittest.main()
