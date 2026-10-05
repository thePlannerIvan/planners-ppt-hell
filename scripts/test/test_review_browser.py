"""Optional browser/interface smoke: the real page driven through the review surface (no plugin host).

要证明的事（B8.3 ＋ B8.4c 的验收）：
  ① 页面加载时注入点已被换成桥的地址（页面里不带桥的副本）；桥声明了 `asset-upload`，控件按它显示；
  ② 一处框选 + 一次逐页提交 → feedback.json 落盘（含上传件）**且** wake 被调用（带 page_key）；
  ③ 陈旧提交被拒：模型重出那一页之后再提交 → 页面提示「这一页已更新，先复核」，不写文件；
  ④ 整套提交用**整句覆盖**唤醒（`{unit:'整套', text:…}`），不是 surface 的逐页模板；
  ⑤ 模型重出一页 → 宿主戳 → 页面**自己换图**：不重载、不丢人写的东西、之后提交不 stale。
All approvals below are synthetic test data.
"""
import hashlib
import json
import os
import sys
import unittest
from pathlib import Path
from urllib.parse import unquote, urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_v5 as fixtures  # noqa: E402
from project_state import PNG, REVIEW, SVG, approved, read  # noqa: E402
from ppt_pipeline import check, export, make_review, next_action, resolve  # noqa: E402
from review_feedback import consume  # noqa: E402
from ppt_pipeline import unresolved  # noqa: E402
import review_surface  # noqa: E402
from lib.planners_modules import resolve_module  # noqa: E402
sys.path.insert(0, str(resolve_module('planners-review-core') / 'scripts' / 'lib'))
import review_host  # noqa: E402  （宿主生命周期只有公共模组那一份）


@unittest.skipUnless(os.environ.get('PPT_BROWSER_TESTS') == '1', 'set PPT_BROWSER_TESTS=1 for local browser/interface smoke')
class ReviewBrowserTests(unittest.TestCase):
    def wire(self, page):
        """把页面与宿主之间**真的发生过的请求**记下来（报告里要给出请求/响应）。"""
        self.exchanges = []

        def on_request(request):
            if '/__review/' in request.url:
                parts = urlsplit(request.url)
                self.exchanges.append(['→', request.method, parts.path + ('?' + parts.query if parts.query else ''),
                                       (request.post_data or '')[:200]])

        def on_response(response):
            if '/__review/' in response.url:
                self.exchanges.append(['←', response.status, urlsplit(response.url).path, ''])

        page.on('request', on_request)
        page.on('response', on_response)

    def submit_page(self, page, label):
        """点一次逐页提交并等到它真的走完。

        先把状态条清空：`#message` 里还留着上一页那句话，直接等它会在**下一次提交之前**就返回
        （第一版踩过：omega 那次点击被 race 掉，文件里还是 pending）。
        """
        page.evaluate("() => { document.querySelector('#message').textContent = '' }")
        page.get_by_text(label, exact=True).click()
        # 提示有三档（宿主已核对 / 宿主未核对 / 同一次提交重发），共同点只有"已写入"——等它，
        # 不要等某一档的措辞，否则换一档就会在这里超时（第一版就踩了）。
        page.wait_for_function("document.querySelector('#message').textContent.includes('已写入')")

    def submit_batch(self, page, label='提交已有决定'):
        page.evaluate("() => { document.querySelector('#message').textContent = '' }")
        page.get_by_text(label, exact=True).click()
        page.wait_for_function("() => { const d = document.querySelector('#submitDialog'); return d && !d.open }")

    def posts(self, suffix):
        return [item for item in self.exchanges if item[0] == '→' and item[1] == 'POST' and item[2].startswith(suffix)]

    def test_real_browser_review_through_the_surface(self):
        fixture = fixtures.V5Tests(); fixture.setUp(); root = fixture.root.resolve()
        hosts = []
        try:
            fixture.payload()                       # 生成审阅页 + snapshot + surface
            state = review_host.start_host(root / review_surface.SURFACE_REL); hosts.append(state)
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                page = browser.new_page(viewport={'width': 1440, 'height': 900})
                errors = []; warnings = []
                page.on('pageerror', lambda error: errors.append(str(error)))
                page.on('console', lambda m: warnings.append(m.text) if m.type == 'warning' else None)
                self.wire(page)
                page.goto(state['url'])
                # 先断言"页面真的长出来了"（页卡 + 导航 + 状态），**不依赖握手**：渲染排在 await 之前
                self.assertEqual(page.locator('#rail button').count(), 2, '加载后必须已经长出逐页卡')
                self.assertIn('1 / 2', page.locator('#title').inner_text(), '当前页标题/导航要在')
                self.assertTrue(page.locator('#pageStatus').inner_text().strip(), '状态/进度要在')
                page.wait_for_function("document.querySelector('#bridgeState').textContent.includes('已连接')")

                def img_src(key):
                    return page.evaluate("(k) => document.querySelector('#rail button[data-key=\"' + k + '\"] img').src", key)

                # ① 注入点被宿主换成了桥的地址；页面里没有桥的副本；能力按声明落实
                html = page.content()
                self.assertNotIn('{{REVIEW_BRIDGE}}', html, '宿主 serve 时必须替换注入点')
                self.assertIn('/__review/bridge.js', html)
                self.assertTrue(page.evaluate('() => !!window.ReviewBridge'), '桥必须已经就位')
                self.assertIn('已连接审阅宿主（http）', page.locator('#bridgeState').inner_text())
                # 桥交到页面手里的能力现在可靠（公共件修掉了"对象属性不跟着闭包更新"）
                capabilities = page.evaluate('() => review.capabilities')
                self.assertIn('asset-upload', capabilities, '宿主声明了上传能力，桥必须把它交到页面手里')
                self.assertTrue(page.locator('#addDrop').is_visible(), '有上传能力时上传控件要在')
                self.assertIn('变化戳：还没收到', page.locator('#pulseState').inner_text())
                page.wait_for_function("document.querySelector('#slide').src.length>0")
                self.assertEqual(page.locator('#rail button').count(), 2)

                # 上传走桥：`review.upload(file, rel)`（宿主按 rel 落在 dir 里）
                page.locator('#newFile').set_input_files(str(root / PNG / 'alpha.png'))
                page.wait_for_selector('.asset')
                self.assertTrue(self.posts('/__review/upload'), '上传必须走桥的 upload')
                self.assertIn('_internal/05_review/uploads/alpha/', unquote(self.posts('/__review/upload')[0][2]))

                # ② 框选一处 + 「本页要求修改」= 写整份状态 + wake(unit=page_key)
                page.get_by_text('框选问题', exact=True).click()
                box = page.locator('#preview').bounding_box()
                page.mouse.move(box['x'] + box['width'] * .1, box['y'] + box['height'] * .1)
                page.mouse.down(); page.mouse.move(box['x'] + box['width'] * .4, box['y'] + box['height'] * .4); page.mouse.up()
                page.locator('#regionNote').fill('标题压到了图')
                page.get_by_text('保存标注', exact=True).click()
                self.assertEqual(page.locator('.annotation').count(), 1)
                page.locator('#feedback').fill('顺便把副标题删掉')
                self.assertIn('要求修改', page.locator('#pageStatus').inner_text())
                self.submit_page(page, '本页要求修改')
                self.assertIn('宿主未核对', page.locator('#message').inner_text(),
                              '宿主没核过就必须说出来，不能和"核过"那档同一句话')
                written = read(root / REVIEW / 'feedback.json')
                snapshot = read(root / REVIEW / 'snapshot.json')
                self.assertEqual(written['pages']['alpha']['decision'], 'revise')
                self.assertEqual(written['pages']['alpha']['feedback'], '顺便把副标题删掉')
                self.assertEqual([a['text'] for a in written['pages']['alpha']['annotations']], ['标题压到了图'])
                self.assertEqual(written['provenance']['source'], 'review_page')
                self.assertEqual(written['pages']['alpha']['version'], snapshot['versions']['alpha'])
                uploaded = written['pages']['alpha']['assets']
                self.assertEqual(len(uploaded), 1, '新增的图片要随提交进反馈')
                self.assertTrue(uploaded[0]['path'].startswith('_internal/05_review/uploads/alpha/'))
                self.assertTrue((root / uploaded[0]['path']).is_file(), '上传件真的落在项目里了')
                self.assertEqual(uploaded[0]['sha256'], hashlib.sha256((root / uploaded[0]['path']).read_bytes()).hexdigest())
                self.assertTrue(self.posts('/__review/write'), '提交必须真的写了一次')
                wakes = self.posts('/__review/wake')
                self.assertTrue(wakes, '逐页提交必须唤醒模型')
                self.assertIn('"unit":"alpha"', wakes[0][3].replace(' ', ''))
                wake_lines = [json.loads(line) for line in (root / review_surface.WAKE_LOG_REL).read_text(encoding='utf-8').splitlines()]
                self.assertEqual(wake_lines[-1]['unit'], 'alpha')
                self.assertIn('alpha 要求修改', wake_lines[-1]['text'], '逐页提交的唤醒语要自带决定，不是"已定"')
                self.assertIn('顺便把副标题删掉', wake_lines[-1]['text'], '唤醒语要带上意见的要点')
                self.assertIn('只重出这一页', wake_lines[-1]['text'])
                self.assertEqual(errors, [])
                print('\n[证据 ②] 页面发出的请求：')
                for item in self.exchanges: print('   ', item[0], item[1], item[2], item[3][:120])

                # ③ 模型重出这一页（同一台宿主、同一个 URL）→ 人再提交必须被拒
                before_writes = len(self.posts('/__review/write'))
                # 模型侧的动作：先按意见改这一页（受影响的页必须真的变了），resolve 掉那一条，再重出审阅页
                item = read(root / REVIEW / 'feedback.json')['items'][0]
                fixture.svg('alpha', 'New title'); fixture.seal()
                resolve(root, item['id'], '标题下移 40px，右侧留白')
                make_review(root)
                page.evaluate("() => { document.querySelector('#message').textContent = '' }")
                page.get_by_text('本页要求修改', exact=True).click()
                page.wait_for_function("document.querySelector('#message').textContent.includes('先复核')")
                self.assertIn('这一页已更新', page.locator('#staleNote').inner_text())
                self.assertEqual(len(self.posts('/__review/write')), before_writes, '陈旧页面不得再写反馈')
                kept = read(root / REVIEW / 'feedback.json')['pages']['alpha']
                self.assertEqual(kept['annotations'][0]['text'], '标题压到了图', '文件仍是上一轮那一份，人的字不丢')
                self.assertEqual(kept['feedback'], '顺便把副标题删掉')
                self.assertEqual(len(kept['assets']), 1, '上一轮的新增图片也要留着（复核时要能看见）')
                intake = consume(root)
                self.assertTrue(intake['stale'], 'Skill 侧也认得出它是上一版')
                self.assertTrue(next_action(root)['review_intake']['stale'])
                print('\n[证据 ③] 陈旧提交被拒：',' | '.join(page.locator('#staleNote').inner_text().split('\n')))

                # 复核：重新加载 → 上一轮的意见是**只读上下文**（不预填进任何参与判定的输入）
                page.reload()
                page.wait_for_function("document.querySelector('#bridgeState').textContent.includes('已连接')")
                previous = page.locator('#previousNote').inner_text()
                self.assertIn('上一轮你说的', previous)
                self.assertIn('顺便把副标题删掉', previous, '上一轮的文字要在只读块里')
                self.assertIn('标题压到了图', previous, '上一轮的框选也要在只读块里')
                self.assertIn('图片改动', previous)
                self.assertEqual(page.locator('#feedback').input_value(), '', '不许把上一轮预填进输入框（那是本轮意见的形状）')
                self.assertEqual(page.locator('.annotation').count(), 0, '上一轮的框选不许进本轮的框选列表')
                self.assertEqual(page.locator('.asset').count(), 0, '上一轮的图片改动不许进本轮的图片列表')
                self.assertIn('未处理', page.locator('#pageStatus').inner_text(), '重出之后初始就是待处理，不是"要求修改"')
                # 一次点击就批准：什么都不用清（验收 ①）
                self.submit_page(page, '本页通过')
                self.assertIn('人工批准', page.locator('#pageStatus').inner_text())
                self.assertIn('alpha 已通过', [json.loads(line)['text'] for line in (root / review_surface.WAKE_LOG_REL).read_text(encoding='utf-8').splitlines()][-1],
                              '通过时的唤醒语是"X 已通过。"')
                page.get_by_text('下一页', exact=True).click()
                page.wait_for_function("document.querySelector('#title').textContent.includes('2 / 2')")
                self.submit_page(page, '本页通过')
                self.assertEqual({k: v['decision'] for k, v in read(root / REVIEW / 'feedback.json')['pages'].items()},
                                 {'alpha': 'approved', 'omega': 'approved'}, '两次逐页通过之后两份决定都在文件里')
                self.assertEqual(read(root / REVIEW / 'feedback.json')['items'], [],
                                 '被重出过的那页不再产生重复待办（上一轮那条已经 resolve 过）')   # 验收 ③
                self.assertEqual(unresolved(root), [])
                self.assertTrue(approved(root), '两页都通过之后整套批准成立')

                # ④ 整套提交也唤醒：**整句覆盖**（surface 的模板是逐页口气，套上去会说成"整套 已定；只重出这一页"）
                page.get_by_role('button', name='整套提交').click()   # 「整套提交」也是对话框标题，按角色点按钮
                self.submit_batch(page)
                self.assertIn('宿主未核对', page.locator('#submitMessage').inner_text(),
                              '整套提交的提示同样要分档')
                self.assertTrue([w for w in self.posts('/__review/wake') if '整套已定' in w[3]], '整套提交要带整句覆盖唤醒')
                last = [json.loads(line) for line in (root / review_surface.WAKE_LOG_REL).read_text(encoding='utf-8').splitlines()][-1]
                self.assertEqual(last['text'], '整套已定，可以进入下一步。')
                self.assertEqual(last['unit'], '整套')
                self.assertTrue(approved(root), '整套提交之后批准仍然成立')
                print('\n[证据 ④] 各次 wake 的请求体：')
                for item in self.posts('/__review/wake'): print('   ', item[1], item[2], item[3])

                # ⑤ 模型重出一页 → 宿主戳 → 页面**自己换图**（不重载、不丢人写的东西、之后提交不 stale）
                page.evaluate("() => { window.__frameMarker = 'marker-' + Math.random().toString(16).slice(2) }")
                marker = page.evaluate('() => window.__frameMarker')
                # 人在**另一页**（omega）写一句 + 框一处，**不提交**
                page.get_by_text('下一页', exact=True).click()
                page.wait_for_function("document.querySelector('#title').textContent.includes('2 / 2')")
                page.locator('#feedback').fill('omega 的意见先放着')
                page.get_by_text('框选问题', exact=True).click()
                box = page.locator('#preview').bounding_box()
                page.mouse.move(box['x'] + box['width'] * .2, box['y'] + box['height'] * .2)
                page.mouse.down(); page.mouse.move(box['x'] + box['width'] * .5, box['y'] + box['height'] * .5); page.mouse.up()
                page.locator('#regionNote').fill('omega 的框选')
                page.get_by_text('保存标注', exact=True).click()
                before = {'alpha': img_src('alpha'), 'omega': img_src('omega'), 'slide': page.locator('#slide').get_attribute('src')}
                # 桥的轮询每 2s 采一次版本令牌，**基准是在第一次采样时建立的**（基准之前改文件，
                # 那次变化永远不会被算成"变了"）。所以改文件之前先让基准落定——这条等的是桥的采样
                # 周期，不是页面的问题（第一版就踩在这个时序上）。
                page.wait_for_timeout(2600)
                fixture.svg('alpha', 'Live updated'); fixture.seal(); make_review(root)   # 模型侧只重出 alpha
                page.wait_for_function("(prev) => { const img = document.querySelector('#rail button[data-key=\"alpha\"] img'); return img && img.src && img.src !== prev; }",
                                       arg=before['alpha'], timeout=15000)
                page.wait_for_timeout(300)     # 让别的缩略图把（同一个）缓存地址写回去
                self.assertEqual(page.evaluate('() => window.__frameMarker'), marker, '① 没有整页重载')
                self.assertNotEqual(img_src('alpha'), before['alpha'], '② 被重出的那页换了图')
                self.assertGreater(page.evaluate("() => document.querySelector('#rail button[data-key=\"alpha\"] img').naturalWidth"), 0, '② 新图真的加载出来了')
                self.assertEqual(img_src('omega'), before['omega'], '② 其他页的 src 一个都没变')
                self.assertEqual(page.locator('#slide').get_attribute('src'), before['slide'], '② 当前页（omega）舞台图没动')
                self.assertEqual(page.locator('#feedback').input_value(), 'omega 的意见先放着', '③ 别的页写的意见原样还在')
                self.assertEqual(page.locator('.annotation').count(), 1, '③ 别的页的框选原样还在')
                pulse = page.locator('#pulseState').inner_text()
                self.assertIn('变化戳：', pulse)
                self.assertNotIn('还没收到', pulse, '变化戳要真的到过页面（状态行就是给这件事用的）')
                print('\n[证据 ⑤] 换图链：帧标记还在=', marker,
                      '| alpha 缩略图 src', before['alpha'][-24:], '→', img_src('alpha')[-24:],
                      '| naturalWidth=', page.evaluate("() => document.querySelector('#rail button[data-key=\"alpha\"] img').naturalWidth"),
                      '| omega 的 src 没动=', img_src('omega') == before['omega'],
                      '|', page.locator('#pulseState').inner_text())
                # ④ 换图后立刻对**别的页**提交：不被判 stale（说明 review_id 跟着更新了）
                self.submit_page(page, '本页要求修改')
                self.assertFalse(consume(root)['stale'], '④ 换图之后提交不该被判 stale')
                self.assertFalse(next_action(root)['review_intake']['stale'])
                decisions = read(root / REVIEW / 'feedback.json')['pages']
                self.assertEqual(decisions['omega']['decision'], 'revise')
                self.assertEqual(decisions['alpha']['decision'], 'pending', '被重出的页退回未处理：旧批准不算数')
                # 切回 alpha：图是新的、状态是"待复核"（面板提示 + rail 的 title）
                page.get_by_text('上一页', exact=True).click()
                page.wait_for_function("document.querySelector('#title').textContent.includes('1 / 2')")
                self.assertFalse(page.locator('#updatedNote').is_hidden(), '⑤ 被重出的页有"画面已更新"的提示')
                self.assertIn('待复核', page.locator('#rail button[data-key="alpha"]').get_attribute('title') or '')
                self.assertFalse(approved(root), '被重出的那一页复核之前，整套不再算批准（旧批准不搬到新画面上）')
                # 验收 ②（真实路径）：被换过图的那页回到"待处理" → 整套提交（含批准未处理页）必须覆盖它，不能跳过
                self.assertEqual(read(root / REVIEW / 'feedback.json')['pages']['alpha']['decision'], 'pending')
                page.get_by_role('button', name='整套提交').click()
                self.submit_batch(page, '批准未处理页并提交')
                final = read(root / REVIEW / 'feedback.json')
                self.assertEqual(final['pages']['alpha']['decision'], 'approved', '整套提交要把待处理页一起批掉（它不该被跳过）')
                # 待办里只该有 omega 那条新意见：被重出过的 alpha 不该再生一条（上一轮那句已经 resolve 过）
                self.assertEqual([i['pages'] for i in final['items']], [['omega']])
                self.assertEqual([i['pages'] for i in unresolved(root)], [['omega']])
                self.assertFalse(approved(root), 'omega 还挂着一条新意见，整套当然还没批准')
                # 规则 4：全部批准、没有新意见之后 → 条目为空，next 报可以导出。
                # omega 那条是他**自己刚写的新意见**（不是历史），所以由他清掉再通过 —— 这是"有意见不许批准"
                # 唯一还管得着的场合：人刚写下的意见。
                page.get_by_text('下一页', exact=True).click()
                page.wait_for_function("document.querySelector('#title').textContent.includes('2 / 2')")
                page.locator('#feedback').fill('')
                page.get_by_text('清空框选', exact=True).click()
                self.submit_page(page, '本页通过')
                settled = read(root / REVIEW / 'feedback.json')
                self.assertEqual(settled['items'], [], '全部批准之后不该还有条目')
                self.assertEqual(unresolved(root), [])
                self.assertTrue(approved(root), '整套批准成立')
                self.assertEqual(next_action(root)['state'], 'EXPORT', '整套批准之后 next 报可以导出')
                # 通知核验分档的另一半：宿主回了 verified → 说"已核对"，并把结果记进记录（同一轮，不是第二次提交）
                page.evaluate('''() => { const original = review.wake;
                    review.wake = async (payload) => Object.assign({}, await original(payload),
                      { verified: { where: 'session-log', seq: 7, target: 'agent/inbox/spliced', requestId: 'r-1', equalsWakeText: true } }); }''')
                self.submit_page(page, '本页通过')
                self.assertIn('宿主已在会话日志里核对进入队列', page.locator('#message').inner_text())
                recorded = read(root / REVIEW / 'feedback.json')['provenance'].get('wake')
                self.assertTrue(recorded and recorded['requested'], '这次通知的核验结果要留在记录里')
                self.assertEqual(recorded['verified']['seq'], 7)
                self.assertEqual(recorded['verified']['equalsWakeText'], True)
                self.assertEqual(read(root / REVIEW / 'feedback.json')['pages']['omega']['decision'], 'approved',
                                 '补写核验结果不改变任何决定')
                # 同一次提交被重发（宿主报 duplicate）：照宿主说的讲——"已在路上"，不是失败、也不是又通知了一次
                page.evaluate('''() => { const original = review.wake;
                    review.wake = async (payload) => Object.assign({}, await original(payload),
                      { duplicate: true, state: 'duplicate',
                        notice: '已写入；本次没有新增通知（同一次提交之前已送达）' }); }''')
                self.submit_page(page, '本页通过')
                shown = page.locator('#message').inner_text()
                self.assertIn('本次没有新增通知（同一次提交之前已送达）', shown, '重发时要说"已在路上"')
                self.assertNotIn('并已通知模型', shown, '不许把"没有新增通知"说得更满')
                # 正常提交（不 duplicate）的文案不变：这一条在上面 ② 与 verified 那一步已经断言过
                self.assertEqual(errors, [])
                self.assertFalse([w for w in warnings if '旧拼法' in w], '不该出现"旧拼法"告警：桥以规范拼法 review/changed 为准')
                self.assertFalse([w for w in warnings if 'review-bridge' in w], '桥不该在 console 出声')
                # ② 反面对照：把 connect() 换成**永不返回** → 页面仍须完整画出，并出现"只读"自述
                stalled = browser.new_page(viewport={'width': 1440, 'height': 900})
                stalled.add_init_script('''
                    Object.defineProperty(window, 'ReviewBridge', { configurable: true,
                      set(v) { if (v && v.connect) { v.connect = function () { return new Promise(function () {}); }; } this.__rb = v; },
                      get() { return this.__rb; } });
                ''')
                stalled.goto(state['url'])
                self.assertEqual(stalled.locator('#rail button').count(), 2, '握手不落地也必须有逐页卡')
                self.assertIn('1 / 2', stalled.locator('#title').inner_text())
                self.assertTrue(stalled.locator('#pageStatus').inner_text().strip())
                stalled.wait_for_function("document.querySelector('#bridgeState').textContent.includes('只读')", timeout=10000)
                self.assertIn('只读', stalled.locator('#bridgeState').inner_text())
                self.assertIn('内容可以看', stalled.locator('#message').inner_text())
                self.assertTrue(stalled.locator('#addDrop').is_hidden(), '没有宿主就没有上传能力，控件要藏起来')
                self.assertTrue(stalled.get_by_text('重新加载', exact=True).is_visible(),
                                '永久 ⟳ 出口必须在（就绪、只读、报错三种状态都在，不依赖任何告警出现）')
                print('\n[证据 ⑥] 没有桥/握手不落地：rail=', stalled.locator('#rail button').count(),
                      '| 自述=', stalled.locator('#bridgeState').inner_text())
                # ③ 结构性对照：把渲染退回"挂到 connect 之后"（旧写法）→ 同一场景下**什么都不画**
                generated = (root / '02_visual_review.html').read_text(encoding='utf-8')
                new_boot = generated[generated.index('(async function boot(){'):generated.index('})();', generated.index('(async function boot(){')) + 5]
                old_boot = ("""(async function boot(){
  if(!window.ReviewBridge){
    $('#bridgeState').textContent='没有加载到审阅桥：这个页面必须由审阅宿主打开（DSH 侧栏，或 serve-review.mjs）。';
    note('页面没连上审阅宿主：写了也存不进去','error');
    return;
  }
  try{review=await ReviewBridge.connect()}
  catch(error){$('#bridgeState').textContent='连不上审阅宿主：'+error.message;note('连不上审阅宿主','error');return}
  render();sizePreview();drawPulse();
})();""")
                legacy = root / '02_visual_review_legacy.html'
                legacy.write_text(generated.replace(new_boot, old_boot), encoding='utf-8')
                red = browser.new_page(viewport={'width': 1440, 'height': 900})
                # 同一场景（握手不落地）下的旧写法：宿主照样给这份 HTML 注入桥，所以「没有桥」不是这里的变量，
                # **握手不返回**才是 —— 旧写法把 render() 挂在 connect() 后面，于是整页空白、且零报错。
                red.add_init_script('''
                    Object.defineProperty(window, 'ReviewBridge', { configurable: true,
                      set(v) { if (v && v.connect) { v.connect = function () { return new Promise(function () {}); }; } this.__rb = v; },
                      get() { return this.__rb; } });
                ''')
                red.goto(state['url'].rsplit('/', 1)[0] + '/02_visual_review_legacy.html')
                red.wait_for_timeout(4000)
                self.assertEqual(red.locator('#rail button').count(), 0,
                                 '旧写法（渲染挂在 connect 之后）在握手不落地时必须**什么都不画** —— 这就是它红的样子')
                self.assertEqual(red.locator('#title').inner_text().strip(), '')
                self.assertEqual(red.locator('#pageStatus').inner_text().strip(), '')
                print('[证据 ⑥ 红对照] 旧写法（渲染在 connect 之后）+ 握手不落地：rail=',
                      red.locator('#rail button').count(), '| 标题=', repr(red.locator('#title').inner_text()),
                      '| 状态=', repr(red.locator('#pageStatus').inner_text()), '（零报错地空着）')
                stalled.close(); red.close()
                browser.close()
            output = export(root)
            self.assertEqual(output['page_count'], 2)
            from pptx import Presentation
            prs = Presentation(root / 'final_deck.pptx'); self.assertTrue(any(s.has_text_frame and s.text for slide in prs.slides for s in slide.shapes))
            self.assertEqual(export(root)['pptx_sha256'], output['pptx_sha256'])
            # 真实渲染器 + 只重出该页的缓存行为
            fixture.svg('alpha', 'Newer title')
            results = check(root, ['alpha']); self.assertEqual(results['pages'][0]['status'], 'pass')
            results = check(root, ['alpha']); self.assertEqual(results['pages'][0]['status'], 'unchanged')
            import subprocess
            result = subprocess.run([sys.executable, str(fixtures.S / 'render_svg_png.py'), str(root / SVG), str(root / 'template-preview-smoke')], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(read(root / 'template-preview-smoke/png_manifest.json')['all_valid_size'])
        finally:
            for item in hosts: review_host.stop_host(item)
            fixture.tearDown()

    def test_typing_feedback_does_not_rebuild_the_rail(self):
        """写反馈打字时，左边的缩略图不该被销毁重建（2026-10-05 真人反馈）。

        `#feedback` 的 oninput → `autoState()` 原本调 `drawRail()`：**每敲一个字**，整条轨道
        连同每一张缩略图的重新取回都重来一遍（插件模式下是每页一次桥往返 + 一个新 blob）。
        人在左边看到的就是缩略图一直在动。

        判据要能证伪：先给每个缩略图打上 `data-probe` 标记，再敲一个字 ——
        标记还在＝没重建（新行为）；标记全没了＝被重建了（旧行为）。
        同时确认状态确实更新了（活动页变成 `revise`），别把"什么都不做"当成修好。
        """
        fixture = fixtures.V5Tests(); fixture.setUp(); root = fixture.root.resolve()
        hosts = []
        try:
            fixture.payload()
            state = review_host.start_host(root / review_surface.SURFACE_REL); hosts.append(state)
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                page = browser.new_page(viewport={'width': 1440, 'height': 900})
                page.goto(state['url'])
                page.wait_for_function("document.querySelector('#bridgeState').textContent.includes('已连接')")
                page.wait_for_function("() => document.querySelectorAll('#rail img').length > 0")
                marked = page.evaluate("""() => {
                  const imgs = [...document.querySelectorAll('#rail img')];
                  imgs.forEach((el, i) => { el.dataset.probe = String(i) });
                  // 顺便数一数这次输入到底写了几个轨道按钮的 className：只该有当前页那一个。
                  let writes = 0;
                  const desc = Object.getOwnPropertyDescriptor(Element.prototype, 'className');
                  Object.defineProperty(Element.prototype, 'className', {
                    configurable: true,
                    get() { return desc.get.call(this) },
                    set(v) { if (this.parentElement && this.parentElement.id === 'rail') writes += 1; desc.set.call(this, v) },
                  });
                  const box = document.querySelector('#feedback');
                  box.value = '写一句意见';
                  box.dispatchEvent(new Event('input', { bubbles: true }));
                  const out = {
                    marked: imgs.length,
                    survived: document.querySelectorAll('#rail img[data-probe]').length,
                    writes: writes,
                    active: document.querySelector('#rail button.active').className,
                    total: document.querySelectorAll('#rail button').length,
                  };
                  Object.defineProperty(Element.prototype, 'className', desc);
                  return out;
                }""")
                self.assertGreater(marked['marked'], 0, '轨道里得先有缩略图，否则这条断言什么也没证')
                self.assertEqual(marked['total'], marked['marked'], '缩略图数量应与页数一致')
                self.assertEqual(marked['survived'], marked['marked'],
                                 '打字只该改状态，不该销毁重建缩略图（重建会让整条轨道一直在动）')
                self.assertEqual(marked['writes'], 1,
                                 '一个字符只该写「当前这一页」那一个按钮的类名，不该把整条轨道写一遍')
                self.assertIn('revise', marked['active'], '状态仍要更新：活动页应变成 revise')
        finally:
            for item in hosts: review_host.stop_host(item)
            fixture.tearDown()

    def test_switching_pages_does_not_refetch_every_thumbnail(self):
        """翻页不该把整套缩略图重新向宿主取一遍（同一条浪费，只是触发频率低些）。

        `render()` 每次都整条重建轨道，而重建就要给每个缩略图取一次 URL（插件模式下是每页
        一次桥往返 + 一个新 blob）。资源按「路径 + 版本」记住之后，第二遍应该一次都不问宿主；
        版本变了（模型重出这一页）自然还是会去取 —— 键里带版本，所以不会拿到旧图。
        """
        fixture = fixtures.V5Tests(); fixture.setUp(); root = fixture.root.resolve()
        hosts = []
        try:
            fixture.payload()
            state = review_host.start_host(root / review_surface.SURFACE_REL); hosts.append(state)
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser = p.chromium.launch(headless=True)
                page = browser.new_page(viewport={'width': 1440, 'height': 900})
                page.goto(state['url'])
                page.wait_for_function("document.querySelector('#bridgeState').textContent.includes('已连接')")
                page.wait_for_function("() => [...document.querySelectorAll('#rail img')].every(i => i.src)")
                page.wait_for_timeout(300)          # 首屏该取的都取完了
                counts = page.evaluate("""async () => {
                  let n = 0;
                  const original = review.asset.bind(review);
                  review.asset = (...a) => { n += 1; return original(...a) };
                  const buttons = [...document.querySelectorAll('#rail button')];
                  buttons[1].click();               // 翻到第 2 页 → render → drawRail
                  await new Promise(r => setTimeout(r, 500));
                  const away = n; n = 0;
                  buttons[0].click();               // 翻回第 1 页
                  await new Promise(r => setTimeout(r, 500));
                  return { away: away, back: n, pages: buttons.length };
                }""")
                self.assertGreaterEqual(counts['pages'], 2, '这条断言要至少两页才有意义')
                self.assertEqual(counts['away'], 0, '翻过去时缩略图与舞台图都该用记住的 URL，不再问宿主')
                self.assertEqual(counts['back'], 0, '翻回来时同样不该重新取')
        finally:
            for item in hosts: review_host.stop_host(item)
            fixture.tearDown()

if __name__ == '__main__':
    unittest.main()
